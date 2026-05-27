from datetime import datetime
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt


def extract_first_markdown_table(markdown_text: str):
    lines = markdown_text.splitlines()
    start_idx = None
    table_lines = []

    for idx, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            if start_idx is None:
                start_idx = idx
            table_lines.append(stripped)
        elif start_idx is not None:
            break

    if start_idx is None or len(table_lines) < 2:
        return None

    header = [cell.strip() for cell in table_lines[0].strip("|").split("|")]
    rows = []
    for line in table_lines[1:]:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(set(cell.replace(" ", "")) <= {"-", ":"} for cell in cells):
            continue
        padded = cells + [""] * (len(header) - len(cells))
        rows.append(padded[: len(header)])

    return {"header": header, "rows": rows}


def add_markdown_table(document: Document, table_data: dict):
    table = document.add_table(rows=1, cols=len(table_data["header"]))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for idx, value in enumerate(table_data["header"]):
        header_cells[idx].text = value

    for row in table_data["rows"]:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value


def build_docx(country: str, report_text: str, date_range: str | None = None):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title_run = title.add_run("Event Source Reporter")
    title_run.font.size = Pt(22)

    subtitle = document.add_paragraph()
    subtitle.style = document.styles["Subtitle"]
    subtitle.add_run(f"Analysis for: {country}")

    meta = document.add_paragraph()
    meta.add_run(f"Generated on: {datetime.now().strftime('%B %d, %Y %H:%M')}")
    if date_range:
        meta.add_run(f"\nDataset range: {date_range}")

    lines = report_text.splitlines()
    idx = 0
    while idx < len(lines):
        stripped = lines[idx].strip()

        if not stripped:
            idx += 1
            continue

        if stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
            idx += 1
            continue
        if stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
            idx += 1
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
            idx += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while idx < len(lines):
                candidate = lines[idx].strip()
                if candidate.startswith("|") and candidate.endswith("|"):
                    table_lines.append(candidate)
                    idx += 1
                else:
                    break
            table_data = extract_first_markdown_table("\n".join(table_lines))
            if table_data:
                add_markdown_table(document, table_data)
            continue

        paragraph_lines = [stripped]
        idx += 1
        while idx < len(lines):
            candidate = lines[idx].strip()
            if not candidate or candidate.startswith("#") or (candidate.startswith("|") and candidate.endswith("|")):
                break
            paragraph_lines.append(candidate)
            idx += 1
        document.add_paragraph(" ".join(paragraph_lines))

    return document


def generate_docx_bytes(country: str, report_text: str, date_range: str | None = None) -> bytes:
    document = build_docx(country, report_text, date_range=date_range)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
